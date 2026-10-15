"""
Dual logging: every violation is appended to both an Excel workbook
(for the "Excel Dashboard & Summarization" module, human-browsable) and a
JSON file (easy to re-load programmatically for the PDF report).
"""
import json
import os
import time
import openpyxl
from openpyxl.utils import get_column_letter
from modules import config

_EXCEL_HEADERS = ["Timestamp", "Student ID", "Violation Type", "Detail", "Screenshot Path"]


def _ensure_workbook():
    if os.path.exists(config.EXCEL_LOG_PATH):
        return openpyxl.load_workbook(config.EXCEL_LOG_PATH)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Violations"
    ws.append(_EXCEL_HEADERS)
    for i, header in enumerate(_EXCEL_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(18, len(header) + 4)
    return wb


def log_violation(student_id: str, v_type: str, detail: str, screenshot_path: str = ""):
    row_ts = time.strftime("%Y-%m-%d %H:%M:%S")

    # --- Excel ---
    wb = _ensure_workbook()
    ws = wb["Violations"]
    ws.append([row_ts, student_id, v_type, detail, screenshot_path])
    wb.save(config.EXCEL_LOG_PATH)

    # --- JSON ---
    record = {
        "timestamp": row_ts,
        "student_id": student_id,
        "violation_type": v_type,
        "detail": detail,
        "screenshot_path": screenshot_path,
    }
    existing = []
    if os.path.exists(config.JSON_LOG_PATH):
        try:
            with open(config.JSON_LOG_PATH, "r") as f:
                existing = json.load(f)
        except (json.JSONDecodeError, FileNotFoundError):
            existing = []
    existing.append(record)
    with open(config.JSON_LOG_PATH, "w") as f:
        json.dump(existing, f, indent=2)


def load_session_events(student_id: str = None):
    if not os.path.exists(config.JSON_LOG_PATH):
        return []
    with open(config.JSON_LOG_PATH, "r") as f:
        events = json.load(f)
    if student_id:
        events = [e for e in events if e["student_id"] == student_id]
    return events


def build_excel_summary_sheet():
    """Adds/refreshes a 'Summary' sheet with a violation-type count table,
    which Streamlit/Excel can chart directly."""
    if not os.path.exists(config.EXCEL_LOG_PATH):
        return
    wb = openpyxl.load_workbook(config.EXCEL_LOG_PATH)
    if "Violations" not in wb.sheetnames:
        return
    ws = wb["Violations"]

    counts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[2] is None:
            continue
        v_type = row[2]
        counts[v_type] = counts.get(v_type, 0) + 1

    if "Summary" in wb.sheetnames:
        del wb["Summary"]
    summary = wb.create_sheet("Summary")
    summary.append(["Violation Type", "Count"])
    for v_type, count in sorted(counts.items(), key=lambda x: -x[1]):
        summary.append([v_type, count])

    wb.save(config.EXCEL_LOG_PATH)
    return counts
